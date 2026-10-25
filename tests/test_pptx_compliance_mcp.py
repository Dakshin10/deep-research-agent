import io
import pytest
from pptx import Presentation

from core.guardrails.compliance_gate import ComplianceGate
from core.tools.pptx_generator import PPTXGenerator
from core.tools.mcp_wrapper import ReadOnlyMCPClient


def test_compliance_gate_seeded_banned_claims():
    """Verify ComplianceGate catches global and tenant-specific banned claims and manages rewrite/block actions."""
    tenant_id = "tenant_marketing_corp"
    gate = ComplianceGate(tenant_id=tenant_id)

    # 1. Test text with global and tenant banned claims
    violating_text = (
        "We promise guaranteed 100% ROI on paid marketing spend with zero down-side risk. "
        "Furthermore, our primary rival is an inferior fraud compared to our platform."
    )

    # Evaluation with 1 rewrite loop
    res = gate.evaluate(violating_text, max_rewrites=1)
    assert res.passed is True
    assert res.rewrite_count == 1
    assert "guaranteed 100% ROI" not in res.sanitized_text
    assert "inferior fraud" not in res.sanitized_text
    assert "[Compliant verified metric]" in res.sanitized_text

    # 2. Test persistent un-rewritable high-severity violation (e.g. max_rewrites=0)
    strict_res = gate.evaluate(violating_text, max_rewrites=0)
    assert strict_res.passed is False
    assert strict_res.delivery_blocked is True
    assert strict_res.flagged_for_human_review is True
    rule_ids = [v.rule_id for v in strict_res.violations]
    assert "unsubstantiated_roi" in rule_ids
    assert "competitor_disparagement" in rule_ids


def test_pptx_layout_checks_max_5_bullets_no_empty_placeholders():
    """
    Test Step 7 requirement:
    PPTX deck generated passes layout checks:
    - Maximum 5 bullets per slide
    - Zero empty placeholders
    - Valid titles on all slides
    """
    title = "Executive Competitor Benchmarking: HubSpot vs Salesforce"
    sections = {
        "Executive Summary": [
            "Market expansion reached 21.4% in 2024 across mid-market B2B segments.",
            "Customer acquisition cost decreased by 14% with automation adoption.",
            "Multi-channel attribution remains the top priority for marketing executives.",
        ],
        "Competitor Comparison": [
            "HubSpot delivers fastest time-to-value for teams under 500 employees.",
            "Salesforce offers deepest custom object workflows for large enterprise.",
            "Zoho CRM provides 40% lower TCO for cost-conscious organizations.",
            "Integration ecosystems show highest parity across core REST APIs.",
            "Security and SOC2 compliance verified across all three vendors.",
            "Extra bullet point 6 that must be truncated to honor max 5 bullets rule.",
            "Extra bullet point 7 that must be truncated.",
        ],
        "Strategic Recommendations": [
            "Consolidate marketing automation under unified CRM architecture.",
            "Enforce strict data governance across customer lifecycle touchpoints.",
        ],
    }

    deck_bytes = PPTXGenerator.generate_deck(title=title, sections=sections)
    assert len(deck_bytes) > 1000

    # Parse with python-pptx to verify layout constraints
    prs = Presentation(io.BytesIO(deck_bytes))

    # Verify slide count
    assert len(prs.slides) == 4  # 1 title slide + 3 section slides

    for idx, slide in enumerate(prs.slides):
        # 1. Slide must have non-empty title
        title_shape = slide.shapes.title
        assert title_shape is not None
        assert title_shape.text.strip() != "", f"Slide {idx} missing title text"

        # 2. Check all shapes for empty placeholders
        for shape in slide.shapes:
            if shape.has_text_frame:
                assert shape.text_frame.text.strip() != "", f"Found empty text frame on slide {idx}"

        # 3. Check that no slide exceeds 5 bullets in body
        if idx > 0:  # content slides
            body = slide.placeholders[1]
            paragraphs = [p.text for p in body.text_frame.paragraphs if p.text.strip()]
            assert len(paragraphs) <= 5, f"Slide {idx} exceeded max 5 bullets: {len(paragraphs)}"


@pytest.mark.asyncio
async def test_mcp_read_only_allowlist_enforcement():
    """
    Test Step 7 requirement:
    MCP wrappers enforce strict read-only allowlisting:
    - HubSpot allow-lists only search/read tools, blocks create/update
    - GA4 allow-lists only read reports
    """
    client = ReadOnlyMCPClient(tenant_id="tenant_agency_test")

    # 1. HubSpot Read Operations (Allowed)
    res_deals = await client.call_tool("hubspot", "search_deals", {"query": "renewal"})
    assert res_deals.read_only is True
    assert res_deals.result["status"] == "success"

    res_comp = await client.call_tool("hubspot", "get_company", {"company_id": "101"})
    assert res_comp.result["status"] == "success"

    # 2. HubSpot Write Operations (MUST be blocked)
    with pytest.raises(PermissionError) as exc_info:
        await client.call_tool("hubspot", "create_deal", {"name": "Malicious Deal", "amount": 100000})
    assert "not allow-listed for read-only access" in str(exc_info.value).lower()

    with pytest.raises(PermissionError) as exc_info:
        await client.call_tool("hubspot", "update_company", {"company_id": "101", "name": "Hacked"})
    assert "not allow-listed for read-only access" in str(exc_info.value).lower()

    # 3. GA4 Read Operations (Allowed)
    res_ga4 = await client.call_tool("ga4", "run_report", {"property_id": "12345"})
    assert res_ga4.read_only is True
    assert len(res_ga4.result["rows"]) > 0

    # 4. GA4 Mutation Operations (MUST be blocked)
    with pytest.raises(PermissionError) as exc_info:
        await client.call_tool("ga4", "delete_property", {"property_id": "12345"})
    assert "not allow-listed" in str(exc_info.value).lower()
